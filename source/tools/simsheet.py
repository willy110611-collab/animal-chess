# 把模擬結果寫進 Excel 的「模擬結果」分頁（改完數值、跑完 xlsx2sheet.py 後執行，讓使用者上線前先看預估表現）
# 用法：python3 tools/simsheet.py 遊戲數據.xlsx [場數]
import json, subprocess, sys, os, datetime
from openpyxl import load_workbook
from openpyxl.styles import Font, PatternFill, Alignment
HERE = os.path.dirname(os.path.abspath(__file__))
path = sys.argv[1]; n = sys.argv[2] if len(sys.argv) > 2 else '40'
D = json.loads(subprocess.check_output(['node', os.path.join(HERE, 'sim_report.js'), n]))
wb = load_workbook(path)
if '模擬結果' in wb.sheetnames: del wb['模擬結果']
ws = wb.create_sheet('模擬結果')
F = 'Arial'
ws.cell(1, 1, f'電腦 8 人混戰模擬 {D["games"]} 場（{datetime.date.today()}），回合中位數 {D["median"]}。平均名次越小越強（1～8），前四＝拿到前四名的比例。這頁每次重新產生，不用改。').font = Font(name=F, italic=True, color='7A6A5A', size=10)
hdr = ['代號', '動物', '費用', '出場隊數', '平均名次', '前四 %', '同費平均差']
for i, h in enumerate(hdr, 1):
    c = ws.cell(3, i, h); c.font = Font(name=F, bold=True, color='FFFFFF'); c.fill = PatternFill('solid', start_color='5A3A2A'); c.alignment = Alignment(horizontal='center')
by_cost = {}
for u in D['units']:
    if u['avg'] is not None: by_cost.setdefault(u['cost'], []).append(u['avg'])
rows = sorted(D['units'], key=lambda u: (u['cost'], u['avg'] if u['avg'] is not None else 9))
for r, u in enumerate(rows, 4):
    mean = sum(by_cost.get(u['cost'], [0])) / max(1, len(by_cost.get(u['cost'], [])))
    diff = round(u['avg'] - mean, 2) if u['avg'] is not None else None
    vals = [u['key'], u['name'], u['cost'], u['n'], u['avg'], u['top'], diff]
    for i, v in enumerate(vals, 1):
        c = ws.cell(r, i, v); c.font = Font(name=F)
        if i == 7 and diff is not None and abs(diff) >= 0.35:
            c.fill = PatternFill('solid', start_color='F8C8B8' if diff < 0 else 'C8E0F8')  # 紅＝太強、藍＝太弱
for i, w in enumerate([12, 10, 6, 10, 10, 8, 12], 1): ws.column_dimensions[chr(64 + i)].width = w
wb.save(path); print('已寫入「模擬結果」分頁：', path)
