# 讀使用者改過的「遊戲數據.xlsx」，跟 data.js 比對，把差異寫成 src/sheet.js（只寫有改的部分）
# 用法：python3 tools/xlsx2sheet.py 遊戲數據.xlsx
import json, subprocess, sys, os
from openpyxl import load_workbook

HERE = os.path.dirname(os.path.abspath(__file__))
RAW = json.loads(subprocess.check_output(['node', os.path.join(HERE, 'dump_data.js'), '--raw']))
wb = load_workbook(sys.argv[1], data_only=False)
errs, notes, need_claude = [], [], []
S = {'CFG': {}, 'ANIMALS': {}, 'NEW_ANIMALS': {}, 'TRAITS': {}, 'MATS': {}, 'ITEMS': {}, 'AUG': {}, 'EVOLVE': {}, 'MISC': {}}
TR_BY_NAME = {v['name']: k for k, v in RAW['TRAITS'].items()}
MAT_BY_NAME = {v['name']: k for k, v in RAW['MATS'].items()}
TIER = {'銀色': 1, '金色': 2, '彩色': 3}
STAT_KEYS = ['atk', 'as', 'hp', 'arm', 'mr', 'sm']

def rows(ws):
    hdr_row = next(r for r in (1, 2, 3) if ws.cell(r, 1).value in ('代號', '等級'))
    for r in ws.iter_rows(min_row=hdr_row + 1, values_only=True):
        if r[0] in (None, ''): continue
        yield [None if (isinstance(v, str) and v.strip() == '') else (v.strip() if isinstance(v, str) else v) for v in r]

def num(v, where):
    if v is None: return None
    try:
        f = float(v); return int(f) if f == int(f) else f
    except Exception:
        errs.append(f'{where}：「{v}」不是數字'); return None

def nums(v, where):
    try: return [int(float(x)) for x in str(v).replace('，', ',').split(',') if x.strip() != '']
    except Exception: errs.append(f'{where}：「{v}」要是逗號分開的數字'); return None

# 基本參數
for k, name, v, *_ in rows(wb['基本參數']):
    if k.startswith('AUG_TIER_ODDS_'):
        i = int(k[-1]) - 1; v2 = nums(v, f'基本參數 {name}')
        if v2 is not None:
            if len(v2) != 3 or sum(v2) != 100: errs.append(f'基本參數 {name}：要 3 個數字且加起來 100')
            elif v2 != RAW['CFG']['AUG_TIER_ODDS'][i]:
                S['CFG'].setdefault('AUG_TIER_ODDS', [list(x) for x in RAW['CFG']['AUG_TIER_ODDS']])[i] = v2
        continue
    if k in ('SYNERGY_BONUS', 'PREY_BONUS'):
        n = num(v, f'基本參數 {name}')
        if n is not None and n != RAW[k]: S['MISC'][k] = n
        continue
    if k not in RAW['CFG']: errs.append(f'基本參數：不認得代號 {k}'); continue
    old = RAW['CFG'][k]
    n = nums(v, f'基本參數 {name}') if isinstance(old, list) else num(v, f'基本參數 {name}')
    if n is not None and n != old: S['CFG'][k] = n
# 商店機率
odds = {}
for r in rows(wb['商店機率']):
    lv = str(int(r[0])); vals = [num(x, f'商店機率 Lv{lv}') or 0 for x in r[1:6]]
    if sum(vals) != 100: errs.append(f'商店機率 Lv{lv}：合計 {sum(vals)}，要等於 100')
    odds[lv] = vals
if odds != RAW['CFG']['SHOP_ODDS']: S['CFG']['SHOP_ODDS'] = odds
# 升級與卡池
ws = wb['升級與卡池']; xp, pool = {}, {}
for r in ws.iter_rows(min_row=2, values_only=True):
    if r[0] not in (None, ''): xp[str(int(r[0]))] = num(r[1], f'升級經驗 Lv{r[0]}')
    if len(r) > 4 and r[3] not in (None, ''): pool[str(int(r[3]))] = num(r[4], f'卡池 {r[3]} 費')
if xp != RAW['CFG']['XP_NEED']: S['CFG']['XP_NEED'] = xp
if pool != RAW['CFG']['POOL']: S['CFG']['POOL'] = pool
# 動物
A_COLS = ['hp', 'atk', 'ai', 'rng', 'arm', 'mr', 'mana', 'sm', 'mv']
for r in rows(wb['動物']):
    k = str(r[0]).strip().lower()
    tr = []
    for t in r[3:5]:
        if t is None: continue
        if t not in TR_BY_NAME: errs.append(f'動物 {r[1]}：不認得羈絆「{t}」'); continue
        tr.append(TR_BY_NAME[t])
    rec = {'name': r[1], 'cost': num(r[2], f'動物 {r[1]} 費用'), 'traits': tr, 'role': r[5] or ''}
    for i, c in enumerate(A_COLS): rec[c] = num(r[6 + i], f'動物 {r[1]} {c}')
    rec['skill'] = r[15] or ''; rec['desc'] = r[16] or ''
    evo = r[17] or ''
    if k not in RAW['ANIMALS']:
        S['NEW_ANIMALS'][k] = rec; need_claude.append(f'新動物「{r[1]}」（{k}）：需要 3D 模型、技能效果、攻擊動畫'); continue
    old = RAW['ANIMALS'][k]; diff = {f: v for f, v in rec.items() if v is not None and old.get(f) != v}
    if diff: S['ANIMALS'][k] = diff
    if evo and RAW['EVOLVE'].get(k) != evo: S['EVOLVE'][k] = evo
# 羈絆
for r in rows(wb['羈絆']):
    k = r[0]
    if k not in RAW['TRAITS']: need_claude.append(f'新羈絆「{r[1]}」：效果需要寫程式'); continue
    rec = {'name': r[1], 'tiers': [num(x, f'羈絆 {r[1]} 門檻') for x in r[2:5]], 'val': [num(x, f'羈絆 {r[1]} 數值') for x in r[5:8]], 'desc': [x or '' for x in r[8:11]]}
    old = RAW['TRAITS'][k]; diff = {f: v for f, v in rec.items() if old.get(f) != v}
    if diff: S['TRAITS'][k] = diff
# 材料
for r in rows(wb['材料']):
    k = r[0]
    if k not in RAW['MATS']: need_claude.append(f'新材料「{r[1]}」：需要掉落設定與圖示'); continue
    stat = {s: num(v, f'材料 {r[1]}') for s, v in zip(STAT_KEYS, r[2:8]) if v is not None}
    old = RAW['MATS'][k]; diff = {}
    if old['name'] != r[1]: diff['name'] = r[1]
    if old['stat'] != stat: diff['stat'] = stat
    if diff: S['MATS'][k] = diff
# 裝備
seen_recipe = {}
for r in rows(wb['裝備']):
    k = str(r[0]).strip().lower()
    if r[2] not in MAT_BY_NAME or r[3] not in MAT_BY_NAME: errs.append(f'裝備 {r[1]}：材料要從清單選'); continue
    frm = [MAT_BY_NAME[r[2]], MAT_BY_NAME[r[3]]]; key = '+'.join(sorted(frm))
    if key in seen_recipe: errs.append(f'裝備 {r[1]} 和 {seen_recipe[key]} 用同一組材料'); continue
    seen_recipe[key] = r[1]
    stat = {s: num(v, f'裝備 {r[1]}') for s, v in zip(STAT_KEYS, r[4:10]) if v is not None}
    rec = {'name': r[1], 'from': frm, 'stat': stat, 'desc': r[10] or ''}
    if k not in RAW['ITEMS']:
        S['ITEMS'][k] = rec; need_claude.append(f'新裝備「{r[1]}」：數值已生效；特殊效果「{rec["desc"]}」需要寫程式'); continue
    old = RAW['ITEMS'][k]; diff = {f: v for f, v in rec.items() if old.get(f) != v and not (f == 'from' and sorted(old['from']) == sorted(v))}
    if diff: S['ITEMS'][k] = diff
# 海克斯
old_aug = {a['id']: a for a in RAW['AUGMENTS']}
for r in rows(wb['海克斯']):
    k = r[0]
    if r[1] not in TIER: errs.append(f'海克斯 {r[2]}：階級要選 銀色／金色／彩色'); continue
    rec = {'tier': TIER[r[1]], 'name': r[2], 'desc': r[3] or ''}
    mr = num(r[4], f'海克斯 {r[2]} 最早回合') if len(r) > 4 else None
    if (mr or 0) != (old_aug.get(k, {}).get('minRound') or 0): rec['minRound'] = mr or 0
    if k not in old_aug: need_claude.append(f'新海克斯「{r[2]}」：效果需要寫程式'); continue
    diff = {f: v for f, v in rec.items() if old_aug[k].get(f) != v}
    if diff: S['AUG'][k] = diff

if errs:
    print('有錯誤，沒有寫入：'); [print(' -', e) for e in errs]; sys.exit(1)
S = {k: v for k, v in S.items() if v}
js = '''// ===== 由「遊戲數據.xlsx」產生（tools/xlsx2sheet.py），請不要手改；要改請改 Excel =====
const SHEET = ''' + json.dumps(S, ensure_ascii=False, indent=1) + ''';
(function applySheet(S) {
  if (S.CFG) { Object.assign(CFG, S.CFG); CFG.ROWS = CFG.ROWS_HALF * 2; }
  for (const k in S.ANIMALS || {}) Object.assign(ANIMALS[k], S.ANIMALS[k]);
  for (const k in S.EVOLVE || {}) EVOLVE[k] = S.EVOLVE[k];
  for (const k in S.TRAITS || {}) Object.assign(TRAITS[k], S.TRAITS[k]);
  for (const k in S.MATS || {}) Object.assign(MATS[k], S.MATS[k]);
  for (const k in S.ITEMS || {}) ITEMS[k] = Object.assign(ITEMS[k] || {}, S.ITEMS[k]);
  for (const k in S.AUG || {}) if (AUG[k]) Object.assign(AUG[k], S.AUG[k]);
  if (S.MISC) { if (S.MISC.SYNERGY_BONUS !== undefined) SYNERGY_BONUS = S.MISC.SYNERGY_BONUS; if (S.MISC.PREY_BONUS !== undefined) PREY_BONUS = S.MISC.PREY_BONUS; }
  for (const k in RECIPE) delete RECIPE[k];
  for (const k in ITEMS) RECIPE[ITEMS[k].from.slice().sort().join('+')] = k;
  for (const t in TRAITS) if (EMBLEMS['emb_' + t]) { EMBLEMS['emb_' + t].name = TRAITS[t].name + '紋章'; EMBLEMS['emb_' + t].desc = '裝上後這隻動物多算一個「' + TRAITS[t].name + '」羈絆'; }
})(SHEET);
'''
open(os.path.join(HERE, '../src/sheet.js'), 'w').write(js)
changed = sum(len(v) for k, v in S.items() if k != 'NEW_ANIMALS')
print(f'已寫入 src/sheet.js：{changed} 項改動')
for k, v in S.items():
    if k == 'NEW_ANIMALS': continue
    for kk, vv in v.items(): print(f'  {k}.{kk} = {json.dumps(vv, ensure_ascii=False)[:120]}')
if need_claude:
    print('需要 Claude 另外處理：'); [print(' -', x) for x in need_claude]
