# 由目前的遊戲資料（data.js＋sheet.js）產生「遊戲數據.xlsx」
# 用法：python3 tools/data2xlsx.py 輸出路徑.xlsx
import json, subprocess, sys, os
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.utils import get_column_letter

HERE = os.path.dirname(os.path.abspath(__file__))
D = json.loads(subprocess.check_output(['node', os.path.join(HERE, 'dump_data.js')]))
out = sys.argv[1] if len(sys.argv) > 1 else 'game_data.xlsx'

F = 'Arial'
HEAD = PatternFill('solid', start_color='5A3A2A')
LOCK = PatternFill('solid', start_color='E7E1DA')
EDIT = PatternFill('solid', start_color='FFF8E1')
THIN = Side(style='thin', color='D9CBB8')
BOX = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)
STAT_KEYS = [('atk', '攻擊'), ('as', '攻速%'), ('hp', '生命'), ('arm', '護甲'), ('mr', '魔抗'), ('sm', '起始能量')]
TR_NAME = {k: v['name'] for k, v in D['TRAITS'].items()}
MAT_NAME = {k: v['name'] for k, v in D['MATS'].items()}
TIER_NAME = {1: '銀色', 2: '金色', 3: '彩色'}

wb = Workbook()

def sheet(title, headers, rows, widths, lock_cols=(0,), wrap_cols=(), note=None):
    ws = wb.create_sheet(title)
    r0 = 1
    if note:
        ws.cell(1, 1, note).font = Font(name=F, italic=True, color='7A6A5A', size=10)
        ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=len(headers))
        r0 = 2
    for j, h in enumerate(headers):
        c = ws.cell(r0, j + 1, h)
        c.font = Font(name=F, bold=True, color='FFFFFF'); c.fill = HEAD; c.alignment = Alignment(horizontal='center', vertical='center', wrap_text=True); c.border = BOX
    for i, row in enumerate(rows):
        for j, v in enumerate(row):
            c = ws.cell(r0 + 1 + i, j + 1, v)
            locked = j in lock_cols
            c.font = Font(name=F, color='7A6A5A' if locked else '0000FF')
            c.fill = LOCK if locked else EDIT
            c.border = BOX
            c.alignment = Alignment(vertical='top', wrap_text=j in wrap_cols)
    for j, w in enumerate(widths):
        ws.column_dimensions[get_column_letter(j + 1)].width = w
    ws.freeze_panes = ws.cell(r0 + 1, 2)
    return ws, r0

def add_list(ws, values, rng):
    dv = DataValidation(type='list', formula1='"' + ','.join(values) + '"', allow_blank=True)
    dv.error = '請從清單選擇'; dv.errorTitle = '不在清單內'
    ws.add_data_validation(dv); dv.add(rng)

# ---- 說明 ----
ws = wb.active; ws.title = '說明'
lines = [
    ('動物戰棋 遊戲數據表', True),
    ('', False),
    ('怎麼用', True),
    ('1. 只改「淡黃底、藍字」的格子；灰底的「代號」是程式用的名字，不要改。', False),
    ('2. 改完存檔，把這個檔案傳給 Claude，說「照 Excel 更新」。Claude 會讀取、跑測試、更新網站與伺服器。', False),
    ('3. 數值改完會先跑模擬檢查平衡，有異常（例如某隻勝率暴衝）會先回報。', False),
    ('', False),
    ('可以新增嗎？', True),
    ('• 新增動物：在「動物」表最下面加一列，代號用英文小寫（例如 panda），其他欄照填。Claude 會補 3D 模型、技能效果與攻擊動畫。', False),
    ('• 新增裝備：在「裝備」表加一列，填兩個材料與數值。數值馬上生效；「說明」裡的特殊效果需要 Claude 寫程式。', False),
    ('• 新增海克斯：在「海克斯」表加一列。效果需要 Claude 寫程式。', False),
    ('• 刪除：不要直接刪列；想停用某隻動物，把它的費用改成 0（不會出現在商店）。', False),
    ('', False),
    ('欄位說明', True),
    ('攻擊間隔、移動：單位是 tick（1 秒 = 20 tick），數字越小越快。', False),
    ('射程：1 = 近戰，2 以上 = 遠程（格數）。', False),
    ('能量上限：放技能需要的能量；起始能量：開場就有的能量。', False),
    ('攻速%：裝備給的攻速百分比（15 = +15%）。', False),
    ('多個數字的欄位用逗號分開，例如「4,5,5,6」。', False),
]
for i, (t, b) in enumerate(lines):
    c = ws.cell(i + 1, 1, t); c.font = Font(name=F, bold=b, size=14 if i == 0 else 11, color='5A3A2A' if b else '333333')
ws.column_dimensions['A'].width = 110
ws.cell(len(lines) + 2, 1, '圖例：').font = Font(name=F, bold=True)
c = ws.cell(len(lines) + 3, 1, '可以改的格子（淡黃底、藍字）'); c.fill = EDIT; c.font = Font(name=F, color='0000FF')
c = ws.cell(len(lines) + 4, 1, '不要改的格子（灰底）'); c.fill = LOCK; c.font = Font(name=F, color='7A6A5A')

# ---- 基本參數 ----
CFG_LABEL = [
    ('COLS', '棋盤欄數', '每邊的欄數（目前 9）'), ('ROWS_HALF', '每邊列數', '每隊半場幾列（目前 5）；總列數自動 ×2'),
    ('PREP_SEC', '準備秒數', '每回合準備階段秒數'), ('COMBAT_MAX_TICK', '戰鬥最長(tick)', '20 tick = 1 秒；700 = 35 秒'),
    ('OVERTIME_TICK', '加時開始(tick)', '幾 tick 後進入加時'), ('NIGHT_TICK', '入夜時間(tick)', '夜行羈絆生效時間'),
    ('TEAM_HP', '團隊血量', '每隊開局血量'), ('BENCH', '備戰席格數', ''), ('MAX_ITEMS', '每隻裝備上限', ''),
    ('START_GOLD', '開局金幣', ''), ('START_LEVEL', '開局等級', ''), ('BASE_INCOME', '基本收入', '（實際依下面「各階段收入」）'),
    ('WIN_GOLD', '贏了多給金幣', ''), ('MAX_INTEREST', '利息上限', '每 10 金 +1，最多幾金'),
    ('INCOME_BY_STAGE', '各階段收入', '第 0、1、2…階段的基本收入，逗號分開'),
    ('XP_COST', '買經驗花費', ''), ('XP_BUY', '買一次得經驗', ''), ('XP_ROUND', '每回合自動經驗', ''),
    ('MAX_LEVEL', '最高等級', ''), ('SHOP_SIZE', '商店張數', ''), ('REROLL_COST', '刷新花費', ''),
    ('POOL_8P_PCT', '8 人局卡池%', '8 人局卡池數量倍率'), ('CAROUSEL_ROUNDS', '選秀回合', '逗號分開'),
    ('ENCOUNTER_ROUNDS', '遭遇事件回合', '逗號分開'), ('AUG_ROUNDS', '海克斯回合', '逗號分開'),
    ('BANK_INTEREST_MAX', '隊伍金庫利息上限', ''), ('STAR_MUL', '星級倍率%', '0 星,1 星,2 星,3 星 的生命攻擊倍率'),
    ('GIFT_GOLD_MAX', '送隊友金幣上限', ''), ('MAX_ROUND', '回合上限', '0 = 不限，打到一隊血量歸零'),
    ('DMG_BASE', '輸了扣血：基本', ''), ('DMG_STAR_PCT', '輸了扣血：存活星數%', ''), ('DMG_STAGE', '輸了扣血：每階段加', ''),
]
rows = []
for k, name, desc in CFG_LABEL:
    v = D['CFG'][k]
    rows.append([k, name, ','.join(str(x) for x in v) if isinstance(v, list) else v, desc])
for i, odds in enumerate(D['CFG']['AUG_TIER_ODDS']):
    rows.append([f'AUG_TIER_ODDS_{i+1}', f'第 {i+1} 次海克斯 銀,金,彩 機率%', ','.join(str(x) for x in odds), '三個數字加起來 100；彩色太早出現會破壞平衡'])
rows.append(['SYNERGY_BONUS', '默契加成%', D['SYNERGY_BONUS'], '兩位隊友都有出同一羈絆時效果加成'])
rows.append(['PREY_BONUS', '天敵加成%', D['PREY_BONUS'], '天敵對獵物的傷害加成'])
sheet('基本參數', ['代號', '名稱', '數值', '說明'], rows, [20, 22, 22, 50], lock_cols=(0, 1, 3))

# ---- 動物 ----
A_COLS = [('hp', '生命'), ('atk', '攻擊'), ('ai', '攻擊間隔'), ('rng', '射程'), ('arm', '護甲'), ('mr', '魔抗'), ('mana', '能量上限'), ('sm', '起始能量'), ('mv', '移動')]
rows = []
for k, a in D['ANIMALS'].items():
    tr = [TR_NAME[t] for t in a['traits']] + ['', '', '']
    rows.append([k, a['name'], a['cost'], tr[0], tr[1], tr[2], a.get('role', '')] + [a[c] for c, _ in A_COLS] + [a.get('skill', ''), a.get('desc', ''), D['EVOLVE'].get(k, '')])
headers = ['代號', '名稱', '費用', '羈絆1', '羈絆2', '羈絆3', '定位'] + [n for _, n in A_COLS] + ['技能名稱', '技能說明', '3 星進化名稱']
ws, r0 = sheet('動物', headers, rows, [12, 12, 7, 9, 9, 9, 10] + [8] * len(A_COLS) + [12, 70, 14], wrap_cols=(17,))
n = len(rows) + 60
add_list(ws, list(TR_NAME.values()), f'D{r0+1}:F{r0+n}')
dv = DataValidation(type='whole', operator='between', formula1='0', formula2='5'); dv.error = '費用要 0～5'; ws.add_data_validation(dv); dv.add(f'C{r0+1}:C{r0+n}')

# ---- 羈絆 ----
rows = []
for k, t in D['TRAITS'].items():
    pad = lambda L: list(L) + [''] * (4 - len(L))
    rows.append([k, t['name']] + pad(t['tiers']) + pad(t['val']) + pad(t['desc']))
sheet('羈絆', ['代號', '名稱', '門檻1', '門檻2', '門檻3', '門檻4', '數值1', '數值2', '數值3', '數值4', '說明1', '說明2', '說明3', '說明4'], rows,
      [10, 8, 7, 7, 7, 7, 7, 7, 7, 7, 30, 24, 24, 30], wrap_cols=(10, 11, 12, 13), note='門檻＝需要幾隻（最多 4 階，沒有第 4 階就留空）；數值＝效果強度（意義見說明）')

# ---- 材料 ----
rows = []
for k, m in D['MATS'].items():
    rows.append([k, m['name']] + [m['stat'].get(s, '') for s, _ in STAT_KEYS])
sheet('材料', ['代號', '名稱'] + [n for _, n in STAT_KEYS], rows, [10, 10] + [9] * len(STAT_KEYS))

# ---- 裝備 ----
rows = []
for k, it in D['ITEMS'].items():
    rows.append([k, it['name'], MAT_NAME[it['from'][0]], MAT_NAME[it['from'][1]]] + [it['stat'].get(s, '') for s, _ in STAT_KEYS] + [it.get('desc', '')])
ws, r0 = sheet('裝備', ['代號', '名稱', '材料1', '材料2'] + [n for _, n in STAT_KEYS] + ['特殊效果說明'], rows,
               [12, 10, 8, 8] + [8] * len(STAT_KEYS) + [56], wrap_cols=(10,), note='兩個材料合成一件；同一組材料只能對應一件裝備')
add_list(ws, list(MAT_NAME.values()), f'C{r0+1}:D{r0+len(rows)+40}')

# ---- 海克斯 ----
rows = [[a['id'], TIER_NAME[a['tier']], a['name'], a['desc'], a.get('minRound', '')] for a in D['AUGMENTS']]
ws, r0 = sheet('海克斯', ['代號', '階級', '名稱', '說明', '最早回合'], rows, [14, 8, 14, 80, 10], wrap_cols=(3,), note='最早回合：這個回合（含）之後才可能出現；空白＝不限。海克斯出現在第 4、9、14 回合')
add_list(ws, list(TIER_NAME.values()), f'B{r0+1}:B{r0+len(rows)+40}')

# ---- 商店機率 ----
rows = [[int(l)] + v for l, v in D['CFG']['SHOP_ODDS'].items()]
ws, r0 = sheet('商店機率', ['等級', '1 費%', '2 費%', '3 費%', '4 費%', '5 費%', '合計'], rows, [8, 9, 9, 9, 9, 9, 9],
               note='每一列合計要等於 100（最右欄自動計算）')
for i in range(len(rows)):
    r = r0 + 1 + i
    c = ws.cell(r, 7, f'=SUM(B{r}:F{r})'); c.font = Font(name=F, color='000000', bold=True); c.fill = LOCK; c.border = BOX

# ---- 升級與卡池 ----
ws = wb.create_sheet('升級與卡池')
for j, h in enumerate(['等級', '升到下一級所需經驗']):
    c = ws.cell(1, j + 1, h); c.font = Font(name=F, bold=True, color='FFFFFF'); c.fill = HEAD; c.border = BOX
for i, (l, v) in enumerate(D['CFG']['XP_NEED'].items()):
    for j, val in enumerate([int(l), v]):
        c = ws.cell(i + 2, j + 1, val); c.font = Font(name=F, color='7A6A5A' if j == 0 else '0000FF'); c.fill = LOCK if j == 0 else EDIT; c.border = BOX
for j, h in enumerate(['費用', '每隻動物張數（全場共用卡池）']):
    c = ws.cell(1, j + 4, h); c.font = Font(name=F, bold=True, color='FFFFFF'); c.fill = HEAD; c.border = BOX
for i, (cst, v) in enumerate(D['CFG']['POOL'].items()):
    for j, val in enumerate([int(cst), v]):
        c = ws.cell(i + 2, j + 4, val); c.font = Font(name=F, color='7A6A5A' if j == 0 else '0000FF'); c.fill = LOCK if j == 0 else EDIT; c.border = BOX
for col, w in zip('ABCDE', [8, 20, 3, 8, 28]): ws.column_dimensions[col].width = w

wb.save(out)
print('saved', out)
